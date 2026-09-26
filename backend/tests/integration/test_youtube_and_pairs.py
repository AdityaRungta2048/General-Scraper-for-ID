"""YouTube adapter + any-platform pairs + user-supplied credentials (setup screen API)."""

from __future__ import annotations

import os
import time
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.main import create_app
from app.models import ProcessingJob
from app.platforms.errors import AuthenticationError, RateLimitedError
from app.platforms.youtube import YouTubeAdapter
from app.services import context as ctxmod
from app.services.jobs import create_job, start_job
from app.workers.processor import process_job
from app.workers.queue import wait_inline
from tests.excel_helpers import make_workbook
from tests.scenarios import MIXED_CASE_ID, img, img_v
from tests.scenarios import build_youtube_world as build_world

TW_YT = ["id_twitch", "country", "id_youtube", "remarks"]


async def run(sf, path: Path, source=None, targets=None) -> ProcessingJob:
    settings = get_settings()
    with sf() as s:
        job, _ = create_job(s, settings, path.name, path.read_bytes())
        start_job(s, job, source, targets)
        job_id = job.id
    await process_job(job_id, settings=settings, sf=sf)
    with sf() as s:
        return s.get(ProcessingJob, job_id)


async def test_twitch_to_youtube_workbook(fake, sf, tmp_path):
    build_world(fake)
    rows = [
        ["starwraith", "Spain", None, None],
        ["linkedalt", "France", None, None],
        ["twitchonlyguy", "Italy", None, None],
        ["ghostie", "Germany", None, None],
        ["nobodyatall", "Spain", None, None],
    ]
    job = await run(sf, make_workbook(tmp_path / "tw_yt.xlsx", TW_YT, rows))
    assert job.status == "COMPLETED", job.error_message
    assert job.verification_json["ok"]
    ws = load_workbook(job.output_path)["Streamers"]
    got = [(ws.cell(row=r, column=3).value, ws.cell(row=r, column=4).value) for r in range(2, 7)]
    assert got == [
        ("@starwraith", None),
        ("@completelydifferent", None),
        (None, "no youtube id found"),
        (None, "no twitch id found"),
        (None, "no Id on both platforms"),
    ]
    assert [ws.cell(row=1, column=c).value for c in (5, 6)] == ["twitch_id_link", "youtube_id_link"]
    assert ws["F2"].value == "https://www.youtube.com/@starwraith"
    assert ws["F2"].hyperlink.target == "https://www.youtube.com/@starwraith"
    assert ws["F5"].value == "https://www.youtube.com/@ghostie"  # unverified same-name suggestion
    assert ws["E2"].value == "https://www.twitch.tv/starwraith"


async def test_youtube_source_accepts_urls_handles_and_ids(fake, sf, tmp_path):
    build_world(fake)
    headers = ["id_youtube", "country", "id_twitch", "remarks"]
    rows = [
        ["https://www.youtube.com/@StarWraith", "Spain", None, None],
        [f"youtube.com/channel/{MIXED_CASE_ID}", "France", None, None],
        ["@ghostie", "Germany", None, None],
    ]
    job = await run(sf, make_workbook(tmp_path / "yt_tw.xlsx", headers, rows))
    assert job.source_platform == "youtube" and job.verification_json["ok"]
    ws = load_workbook(job.output_path)["Streamers"]
    assert ws["C2"].value == "StarWraith"
    assert ws["C3"].value == "LinkedAlt"
    assert (ws["C4"].value, ws["D4"].value) == (None, "no twitch id found")


async def test_youtube_probing_is_capped_to_save_quota(fake, sf, tmp_path):
    fake.add_twitch("somebodyunique", "SomebodyUnique", image=img(9))
    await run(sf, make_workbook(tmp_path / "q.xlsx", TW_YT, [["somebodyunique", "Spain", None, None]]))
    by_handle = sum(1 for r in fake.requests if "forHandle" in str(r.url))
    searches = sum(1 for r in fake.requests if r.url.path.endswith("/search"))
    assert by_handle <= YouTubeAdapter.variant_budget
    assert searches <= 1


async def test_youtube_quota_exhaustion_is_retryable_not_no_match(fake, sf, tmp_path):
    build_world(fake)

    def handle(request):
        if request.url.host == "www.googleapis.com":  # Google's "daily quota used up" answer
            return httpx.Response(403, json={"error": {"code": 403, "errors": [{"reason": "quotaExceeded"}]}})
        return fake.handle(request)

    ctxmod.set_default_transport(httpx.MockTransport(handle))
    job = await run(sf, make_workbook(tmp_path / "quota.xlsx", TW_YT, [["starwraith", "Spain", None, None]]))
    assert job.status == "PARTIAL" and job.error_count == 1
    ws = load_workbook(job.output_path)["Streamers"]
    assert ws["C2"].value is None and ws["D2"].value is None  # row untouched, retry later


def test_normalize_youtube_values():
    n = YouTubeAdapter.normalize_handle
    assert n("@MrBeast") == "mrbeast"
    assert n("https://www.youtube.com/@MrBeast/videos") == "mrbeast"
    assert n(MIXED_CASE_ID) == "channel:" + MIXED_CASE_ID
    assert n(f"https://youtube.com/channel/{MIXED_CASE_ID}") == "channel:" + MIXED_CASE_ID
    assert n("https://twitch.tv/foo") is None
    assert n("ab") is None


# ------------------------------------------------------------------ setup screen API
@pytest.fixture
def client(fake):
    with TestClient(create_app()) as c:
        yield c


def test_platform_status_never_returns_key_values(client):
    rows = client.get("/api/platforms").json()
    assert {r["id"] for r in rows} == {
        "twitch", "kick", "youtube", "chzzk", "soop", "steam", "bigo", "nimo", "rumble", "brave",
    }  # fmt: skip
    assert "test-yt-key" not in str(rows) and "test-twitch-secret" not in str(rows)
    yt = next(r for r in rows if r["id"] == "youtube")
    assert yt["fields"] == ["api_key"] and yt["configured"] and yt["source"] == "environment"


def test_saving_credentials_checks_them_first(client, monkeypatch):
    monkeypatch.setenv("YOUTUBE_API_KEY", "")
    get_settings.cache_clear()
    assert not next(r for r in client.get("/api/platforms").json() if r["id"] == "youtube")["configured"]

    bad = client.put("/api/credentials/youtube", json={"values": {"api_key": "wrong"}})
    assert bad.status_code == 400 and "rejected" in bad.json()["detail"]
    missing = client.put("/api/credentials/youtube", json={"values": {"api_key": "  "}})
    assert missing.status_code == 422

    ok = client.put("/api/credentials/youtube", json={"values": {"api_key": "test-yt-key"}})
    assert ok.status_code == 200 and ok.json()["configured"] and ok.json()["source"] == "setup"
    path = get_settings().data_dir / "credentials.json"
    assert path.exists() and (os.name == "nt" or (path.stat().st_mode & 0o077) == 0)

    cleared = client.delete("/api/credentials/youtube").json()
    assert not cleared["configured"]


def test_starting_a_job_without_keys_says_what_is_missing(client, tmp_path, monkeypatch):
    monkeypatch.setenv("YOUTUBE_API_KEY", "")
    get_settings.cache_clear()
    path = make_workbook(tmp_path / "x.xlsx", TW_YT, [["starwraith", "Spain", None, None]])
    with path.open("rb") as fh:
        job = client.post("/api/jobs", files={"file": ("x.xlsx", fh)}).json()
    r = client.post(f"/api/jobs/{job['id']}/start", json={})
    assert r.status_code == 409 and "YouTube" in r.json()["detail"]


def _finish(client, job_id):
    wait_inline(job_id, timeout=60)
    for _ in range(100):
        done = client.get(f"/api/jobs/{job_id}").json()
        if done["status"] in ("COMPLETED", "PARTIAL", "FAILED"):
            return done
        time.sleep(0.05)
    raise AssertionError("job did not finish")


def build_multi_world(fake):
    build_world(fake)
    fake.add_kick("starwraith", "StarWraith", image=img_v(1), description="twitch.tv/starwraith")


MULTI_ROWS = ["starwraith", "linkedalt", "twitchonlyguy", "ghostie", "nobodyatall"]


def test_one_source_searched_on_several_platforms_at_once(client, fake, tmp_path):
    build_multi_world(fake)
    # only the source column exists: the destination id columns are added to the output
    path = make_workbook(
        tmp_path / "multi.xlsx", ["id_twitch", "country", "remarks"], [[v, "Spain", None] for v in MULTI_ROWS]
    )
    with path.open("rb") as fh:
        job = client.post("/api/jobs", files={"file": ("multi.xlsx", fh)}).json()
    assert job["detected_platform"] == "twitch" and job["platforms"] == ["twitch"]
    assert client.post(f"/api/jobs/{job['id']}/start", json={}).status_code == 409  # no destination yet
    r = client.post(
        f"/api/jobs/{job['id']}/start",
        json={"source_platform": "twitch", "target_platforms": ["kick", "youtube"]},
    )
    assert r.status_code == 200, r.text
    done = _finish(client, job["id"])
    assert done["status"] == "COMPLETED", done["error_message"]
    assert done["target_platforms"] == ["kick", "youtube"] and done["total_rows"] == 10
    assert done["verification_json"]["ok"]

    ws = load_workbook(get_settings().outputs_dir / job["id"] / "multi_processed.xlsx")["Streamers"]
    header = [ws.cell(row=1, column=c).value for c in range(1, 9)]
    assert header == [
        "id_twitch", "country", "remarks", "id_kick", "id_youtube",
        "twitch_id_link", "kick_id_link", "youtube_id_link",
    ]  # fmt: skip
    got = [
        (ws.cell(row=r, column=4).value, ws.cell(row=r, column=5).value, ws.cell(row=r, column=3).value)
        for r in range(2, 7)
    ]
    assert got == [
        ("starwraith", "@starwraith", None),
        (None, "@completelydifferent", "no kick id"),
        (None, None, "no kick id; no youtube id found"),
        (None, None, "no twitch id found; no kick id"),
        (None, None, "no Id on any platform"),
    ]
    assert (
        ws["G2"].value == "https://kick.com/starwraith"
        and ws["H2"].value == "https://www.youtube.com/@starwraith"
    )
    assert ws["F2"].value == "https://www.twitch.tv/starwraith"

    rows = client.get(f"/api/jobs/{job['id']}/rows").json()["items"]
    assert [(r["original_row"], r["target_platform"]) for r in rows[:2]] == [(2, "kick"), (2, "youtube")]
    detail = client.get(f"/api/jobs/{job['id']}/rows/2?target=youtube").json()
    assert detail["target_platform"] == "youtube" and detail["matched_id"] == "@starwraith"


def test_review_on_one_destination_updates_the_combined_remark(client, fake, tmp_path):
    build_multi_world(fake)
    path = make_workbook(
        tmp_path / "rv.xlsx", ["id_twitch", "country", "remarks"], [["ghostie", "Spain", None]]
    )
    with path.open("rb") as fh:
        job = client.post("/api/jobs", files={"file": ("rv.xlsx", fh)}).json()
    client.post(f"/api/jobs/{job['id']}/start", json={"target_platforms": ["kick", "youtube"]})
    _finish(client, job["id"])
    items = client.get(f"/api/jobs/{job['id']}/reviews").json()
    assert [i["target_platform"] for i in items] == ["youtube"]
    ambiguous = client.post(f"/api/jobs/{job['id']}/reviews", json={"original_row": 2, "verdict": "CONFIRM"})
    assert ambiguous.status_code == 400  # two destinations for this row: say which
    ok = client.post(
        f"/api/jobs/{job['id']}/reviews",
        json={
            "original_row": 2,
            "verdict": "CONFIRM",
            "target_platform": "youtube",
            "target_username": "ghostie",
        },
    )
    assert ok.status_code == 200, ok.text
    assert client.get(f"/api/jobs/{job['id']}/download").status_code == 200
    ws = load_workbook(get_settings().outputs_dir / job["id"] / "rv_processed.xlsx")["Streamers"]
    assert (ws["D2"].value, ws["E2"].value, ws["C2"].value) == (
        None,
        "@ghostie",
        "no twitch id found; no kick id",
    )


def test_one_failed_destination_leaves_the_remark_untouched(client, fake, tmp_path):
    build_multi_world(fake)

    def handle(request):
        if request.url.host == "www.googleapis.com":
            return httpx.Response(403, json={"error": {"errors": [{"reason": "quotaExceeded"}]}})
        return fake.handle(request)

    ctxmod.set_default_transport(httpx.MockTransport(handle))
    path = make_workbook(
        tmp_path / "pf.xlsx", ["id_twitch", "country", "remarks"], [["starwraith", "Spain", None]]
    )
    with path.open("rb") as fh:
        job = client.post("/api/jobs", files={"file": ("pf.xlsx", fh)}).json()
    client.post(f"/api/jobs/{job['id']}/start", json={"target_platforms": ["kick", "youtube"]})
    done = _finish(client, job["id"])
    assert done["status"] == "PARTIAL" and done["error_count"] == 1
    ws = load_workbook(get_settings().outputs_dir / job["id"] / "pf_processed.xlsx")["Streamers"]
    assert (ws["D2"].value, ws["E2"].value, ws["C2"].value) == ("starwraith", None, None)


def test_quota_and_bad_key_errors_are_classified():
    import asyncio

    from app.platforms.http import PlatformHttpClient, RetryPolicy, StaticKeyToken

    def handler(request):
        reason = request.url.params["r"]
        return httpx.Response(
            403 if reason == "quotaExceeded" else 400, json={"error": {"errors": [{"reason": reason}]}}
        )

    async def call(reason):
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as c:
            http = PlatformHttpClient(
                platform="youtube",
                base_url="https://www.googleapis.com/youtube/v3",
                client=c,
                token=StaticKeyToken("youtube", "k"),
                retry=RetryPolicy(0, 0),
            )
            await http.get_json("channels", params={"r": reason})

    with pytest.raises(RateLimitedError):
        asyncio.run(call("quotaExceeded"))
    with pytest.raises(AuthenticationError):
        asyncio.run(call("keyInvalid"))
