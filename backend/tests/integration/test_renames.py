"""Renamed channels, found through the Streams Charts API (optional key)."""

from __future__ import annotations

from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.main import create_app
from app.models import ProcessingJob
from app.services.jobs import create_job, start_job
from app.workers.processor import process_job
from tests.excel_helpers import make_workbook

TW_KICK = ["id_twitch", "country", "id_kick", "remarks"]


@pytest.fixture
def sc_key(env, monkeypatch):
    monkeypatch.setenv("STREAMSCHARTS_CLIENT_ID", "test-sc-id")
    monkeypatch.setenv("STREAMSCHARTS_TOKEN", "test-sc-token")
    get_settings.cache_clear()
    return get_settings()


def build_world(fake) -> None:
    # the sheet's Twitch ID was renamed; the Kick channel has the new name
    fake.add_twitch("oldtwitch", language="de")
    fake.rename("twitch", "oldtwitch", "newtwitch")
    fake.add_kick("newtwitch", language="de")
    # the Kick channel was renamed away from the Twitch name
    fake.add_twitch("kickrenamed")
    fake.add_kick("kickrenamed")
    fake.rename("kick", "kickrenamed", "kickrenamed-es")
    # Streams Charts' id leads to a Kick account that looks nothing like what it reported
    fake.add_twitch("unconfirmed")
    fake.add_kick("unconfirmed")
    fake.rename("kick", "unconfirmed", "otherkick")
    fake.streamscharts[("kick", "unconfirmed")] = {
        **fake.streamscharts[("kick", "unconfirmed")],
        "channel_name": "unconfirmed",
        "channel_display_name": "Unconfirmed",
        "avatar_url": "https://files.kick.com/images/someone-else.png",
    }
    # never on Kick at all
    fake.add_twitch("twitchonly")
    # both renamed: the Twitch ID and the Kick channel that had the same old name
    fake.add_twitch("oldboth")
    fake.rename("twitch", "oldboth", "newboth")
    fake.add_kick("oldboth")
    fake.rename("kick", "oldboth", "newboth-k")


ROWS = [
    ["oldtwitch", "Germany", None, None],
    ["kickrenamed", "Spain", None, None],
    ["unconfirmed", "France", None, None],
    ["twitchonly", "Italy", None, None],
    ["oldboth", "Spain", None, None],
]


async def run(sf, path: Path) -> ProcessingJob:
    settings = get_settings()
    with sf() as s:
        job, _ = create_job(s, settings, path.name, path.read_bytes())
        start_job(s, job, None, None)
        job_id = job.id
    await process_job(job_id, settings=settings, sf=sf)
    with sf() as s:
        return s.get(ProcessingJob, job_id)


def read(job: ProcessingJob) -> list[dict[str, object]]:
    ws = load_workbook(job.output_path)["Streamers"]
    headers = [c.value for c in ws[1]]
    return [
        {h: ws.cell(row=r, column=i + 1).value for i, h in enumerate(headers) if h}
        for r in range(2, ws.max_row + 1)
    ]


async def test_renamed_channels_are_followed(fake, sf, sc_key, tmp_path):
    build_world(fake)
    job = await run(sf, make_workbook(tmp_path / "renames.xlsx", TW_KICK, ROWS))
    assert job.status == "COMPLETED", job.error_message
    assert job.verification_json["ok"]
    cols = ("id_twitch", "id_kick", "remarks", "twitch_id_link", "kick_id_link")
    got = [tuple(r[c] for c in cols) for r in read(job)]
    assert got == [
        # the sheet's ID is kept; the link points to the channel's current name
        ("oldtwitch", "newtwitch", None, "https://www.twitch.tv/newtwitch", "https://kick.com/newtwitch"),
        ("kickrenamed", "kickrenamed-es", None, "https://www.twitch.tv/kickrenamed", "https://kick.com/kickrenamed-es"),
        # an unconfirmed rename goes to Review: nothing written
        ("unconfirmed", None, "no kick id", "https://www.twitch.tv/unconfirmed", None),
        ("twitchonly", None, "no kick id", "https://www.twitch.tv/twitchonly", None),
        ("oldboth", "newboth-k", None, "https://www.twitch.tv/newboth", "https://kick.com/newboth-k"),
    ]  # fmt: skip
    assert job.review_count == 1


async def test_without_a_streams_charts_key_nothing_changes(fake, sf, tmp_path):
    build_world(fake)
    job = await run(sf, make_workbook(tmp_path / "renames.xlsx", TW_KICK, ROWS))
    assert job.status == "COMPLETED", job.error_message
    remarks = [r["remarks"] for r in read(job)]
    both = "no Id on both platforms"
    assert remarks == [both, "no kick id", "no kick id", "no kick id", both]
    assert fake.api_calls("streamscharts.com") == 0


async def test_streams_charts_is_asked_only_when_nothing_matches(fake, sf, sc_key, tmp_path):
    fake.add_twitch("samename")
    fake.add_kick("samename")
    fake.add_twitch("twitchonly")
    rows = [["samename", "Spain", None, None], ["twitchonly", "Spain", None, None]]
    job = await run(sf, make_workbook(tmp_path / "plain.xlsx", TW_KICK, rows))
    assert [r["remarks"] for r in read(job)] == [None, "no kick id"]
    # one (free) 404 for the unmatched row; nothing for the matched one
    assert fake.api_calls("streamscharts.com") == 1


async def test_a_bot_check_is_reported_not_bypassed(fake, sf, sc_key, tmp_path):
    fake.add_twitch("twitchonly")
    fake.fail(lambda r: r.url.host == "streamscharts.com", status=403, headers={"cf-mitigated": "challenge"})
    job = await run(sf, make_workbook(tmp_path / "cf.xlsx", TW_KICK, [["twitchonly", "Spain", None, None]]))
    assert job.status == "FAILED"
    assert "bot check" in job.error_message
    assert fake.api_calls("streamscharts.com") == 1  # asked once, never retried around


def test_saving_the_key_uses_free_testing_mode(fake):
    with TestClient(create_app()) as client:
        url = "/api/credentials/streamscharts"
        bad = client.put(url, json={"values": {"client_id": "test-sc-id", "token": "x"}})
        assert bad.status_code == 400
        ok = client.put(url, json={"values": {"client_id": "test-sc-id", "token": "test-sc-token"}})
        assert ok.status_code == 200 and ok.json()["configured"] and ok.json()["optional"]
        assert "test-sc-token" not in str(client.get("/api/platforms").json())
    checks = [r for r in fake.requests if r.url.host == "streamscharts.com"]
    assert checks and all(r.url.params.get("testing_mode") == "true" for r in checks)
