"""All nine platforms: Twitch, Kick, YouTube, CHZZK, SOOP, Steam, Bigo LIVE, Nimo TV, Rumble."""

from __future__ import annotations

from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from app.config import get_settings
from app.models import ProcessingJob
from app.platforms.catalog import PLATFORMS
from app.services import context as ctxmod
from app.services.jobs import create_job, start_job
from app.workers.processor import process_job
from tests.excel_helpers import make_workbook
from tests.scenarios import build_all_platforms_world as build_everywhere

NEW = ["chzzk", "soop", "steam", "bigo", "nimo", "rumble"]
ALL_TARGETS = ["kick", "youtube", *NEW]


async def run(sf, path: Path, source=None, targets=None) -> ProcessingJob:
    settings = get_settings()
    with sf() as s:
        job, _ = create_job(s, settings, path.name, path.read_bytes())
        start_job(s, job, source, targets)
        job_id = job.id
    await process_job(job_id, settings=settings, sf=sf)
    with sf() as s:
        return s.get(ProcessingJob, job_id)


async def test_twitch_ids_searched_on_all_eight_other_platforms(fake, sf, tmp_path):
    ids = build_everywhere(fake)
    path = make_workbook(
        tmp_path / "all.xlsx",
        ["id_twitch", "country", "remarks"],
        [["starwraith", "Spain", None], ["nobodyhere", "Spain", None]],
    )
    job = await run(sf, path, "twitch", ALL_TARGETS)
    assert job.status == "COMPLETED", job.error_message
    assert job.verification_json["ok"]
    ws = load_workbook(job.output_path)["Streamers"]
    header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    assert header[:3] == ["id_twitch", "country", "remarks"]
    col = {h: i + 1 for i, h in enumerate(header)}
    expected = {
        "kick": "starwraith",
        "youtube": "@starwraith",
        "chzzk": ids["chzzk"],
        "soop": "starwraith",
        "steam": "starwraith",
        "bigo": "starwraith",
        "nimo": "starwraith",
        "rumble": "starwraith",
    }
    got = {p: ws.cell(row=2, column=col[f"id_{p}"]).value for p in ALL_TARGETS}
    assert got == expected
    assert ws.cell(row=2, column=3).value is None  # confirmed everywhere: no remark
    assert ws.cell(row=3, column=3).value == "no Id on any platform"
    assert ws.cell(row=2, column=col["soop_id_link"]).value == "https://ch.sooplive.co.kr/starwraith"
    assert ws.cell(row=2, column=col["chzzk_id_link"]).value == f"https://chzzk.naver.com/{ids['chzzk']}"
    assert ws.cell(row=2, column=col["rumble_id_link"]).value == "https://rumble.com/c/starwraith"


@pytest.mark.parametrize(
    ("platform", "value"),
    [
        ("soop", "https://ch.sooplive.co.kr/starwraith"),
        ("bigo", "https://www.bigo.tv/user/starwraith"),
        ("nimo", "https://www.nimo.tv/starwraith"),
        ("rumble", "https://rumble.com/c/StarWraith"),
        ("steam", "https://steamcommunity.com/id/starwraith/"),
    ],
)
async def test_new_platforms_as_source(fake, sf, tmp_path, platform, value):
    build_everywhere(fake)
    path = make_workbook(
        tmp_path / f"{platform}.xlsx",
        [f"id_{platform}", "country", "id_twitch", "remarks"],
        [[value, "Spain", None, None]],
    )
    job = await run(sf, path)
    assert job.source_platform == platform and job.verification_json["ok"], job.error_message
    ws = load_workbook(job.output_path)["Streamers"]
    assert ws["C2"].value == "StarWraith", ws["D2"].value


async def test_chzzk_source_by_channel_url(fake, sf, tmp_path):
    ids = build_everywhere(fake)
    path = make_workbook(
        tmp_path / "chzzk.xlsx",
        ["id_chzzk", "country", "id_twitch", "remarks"],
        [[f"https://chzzk.naver.com/live/{ids['chzzk']}", "Korea", None, None]],
    )
    job = await run(sf, path)
    ws = load_workbook(job.output_path)["Streamers"]
    assert ws["C2"].value == "StarWraith"


async def test_blocked_public_site_is_a_retryable_row_error_not_a_job_failure(fake, sf, tmp_path):
    build_everywhere(fake)

    def handle(request):
        if request.url.host == "rumble.com":
            return httpx.Response(403, text="blocked")
        return fake.handle(request)

    ctxmod.set_default_transport(httpx.MockTransport(handle))
    path = make_workbook(
        tmp_path / "blocked.xlsx", ["id_twitch", "country", "remarks"], [["starwraith", "Spain", None]]
    )
    job = await run(sf, path, "twitch", ["soop", "rumble"])
    assert job.status == "PARTIAL" and job.error_count == 1
    ws = load_workbook(job.output_path)["Streamers"]
    assert ws["D2"].value == "starwraith"  # soop still written
    assert ws["C2"].value is None  # remark waits for the rumble retry


def test_every_platform_is_in_the_catalog_and_parses_its_urls():
    assert list(PLATFORMS) == [
        "twitch",
        "kick",
        "youtube",
        "chzzk",
        "soop",
        "steam",
        "bigo",
        "nimo",
        "rumble",
    ]
    n = {p: PLATFORMS[p].adapter.normalize_handle for p in PLATFORMS}
    assert n["chzzk"]("https://chzzk.naver.com/live/" + "a" * 32) == "a" * 32
    assert n["chzzk"]("StarWraith") is None  # channel names aren't ids
    assert n["soop"]("https://play.sooplive.co.kr/phonics1/2812") == "phonics1"
    assert n["soop"]("bj.afreecatv.com/phonics1") == "phonics1"
    assert n["steam"]("76561197960287930") == "76561197960287930"
    assert n["steam"]("https://steamcommunity.com/profiles/76561197960287930") == "76561197960287930"
    assert n["bigo"]("https://www.bigo.tv/en/api-docs") == "api-docs"
    assert n["nimo"]("@Someone") == "someone"
    assert n["rumble"]("https://rumble.com/user/SomeOne") == "someone"
    assert n["rumble"]("https://twitch.tv/x") is None
    assert {p for p, i in PLATFORMS.items() if i.keyless} == {"soop", "bigo", "nimo", "rumble"}


@pytest.fixture
def client(fake):
    from fastapi.testclient import TestClient

    from app.main import create_app

    with TestClient(create_app()) as c:
        yield c


def test_keyless_platforms_are_switched_on_not_given_keys(client, monkeypatch, tmp_path):
    monkeypatch.setenv("PUBLIC_LOOKUP_PLATFORMS", "")
    get_settings.cache_clear()
    status = {p["id"]: p for p in client.get("/api/platforms").json()}
    assert (
        status["rumble"]["keyless"]
        and status["rumble"]["fields"] == []
        and not status["rumble"]["configured"]
    )
    assert not status["chzzk"]["keyless"] and status["chzzk"]["fields"] == ["client_id", "client_secret"]

    path = make_workbook(
        tmp_path / "r.xlsx", ["id_twitch", "country", "remarks"], [["starwraith", "Spain", None]]
    )
    with path.open("rb") as fh:
        job = client.post("/api/jobs", files={"file": ("r.xlsx", fh)}).json()
    blocked = client.post(f"/api/jobs/{job['id']}/start", json={"target_platforms": ["rumble"]})
    assert blocked.status_code == 409 and "Rumble" in blocked.json()["detail"]

    on = client.put("/api/credentials/rumble", json={"values": {}}).json()
    assert on["configured"] and on["source"] == "setup"
    assert (
        client.post(f"/api/jobs/{job['id']}/start", json={"target_platforms": ["rumble"]}).status_code == 200
    )
    assert not client.delete("/api/credentials/rumble").json()["configured"]


@pytest.mark.parametrize(
    ("cred", "good", "bad"),
    [
        (
            "chzzk",
            {"client_id": "test-chzzk-id", "client_secret": "test-chzzk-secret"},
            {"client_id": "x", "client_secret": "y"},
        ),
        ("steam", {"api_key": "test-steam-key"}, {"api_key": "wrong"}),
    ],
)
def test_chzzk_and_steam_keys_are_checked_before_saving(client, cred, good, bad):
    assert client.put(f"/api/credentials/{cred}", json={"values": bad}).status_code == 400
    ok = client.put(f"/api/credentials/{cred}", json={"values": good})
    assert ok.status_code == 200 and ok.json()["configured"]
