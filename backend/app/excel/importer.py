"""ExcelImporter: validation, header location, source-platform detection, row extraction.

Each data row keeps its immutable 1-based Excel row number (``original_row``).
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.excel.state_machine import is_blank
from app.platforms.catalog import PLATFORMS, label

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}
HEADER_SCAN_ROWS = 20
MAX_DATA_ROWS = 200_000

_ALIASES: dict[str, str] = {
    "country": "country",
    "pays": "country",
    "land": "country",
    "pais": "country",
    "paese": "country",
    "remarks": "remarks",
    "remark": "remarks",
    "comments": "remarks",
    "comment": "remarks",
    "notes": "remarks",
    "yt": "id_youtube",
    "id_yt": "id_youtube",
    "yt_id": "id_youtube",
    "yt_id_link": "youtube_id_link",
}
for _p in PLATFORMS:
    _ALIASES.update(dict.fromkeys((f"id_{_p}", f"{_p}_id", f"id{_p}", _p), f"id_{_p}"))
    _ALIASES.update(
        dict.fromkeys((f"{_p}_id_link", f"{_p}_link", f"{_p}_url", f"{_p}_channel_link"), f"{_p}_id_link")
    )
ID_HEADERS = {f"id_{p}": p for p in PLATFORMS}
LINK_HEADERS = {p: f"{p}_id_link" for p in PLATFORMS}


class WorkbookValidationError(ValueError):
    pass


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = re.sub(r"[\s\-\.]+", "_", text)
    text = re.sub(r"_+", "_", text).strip("_")
    return _ALIASES.get(text, text)


@dataclass
class ColumnMap:
    sheet_name: str
    header_row: int
    ids: dict[str, int]  # platform -> id column, e.g. {"kick": 1, "twitch": 3}
    country: int | None
    remarks: int | None  # None => the exporter adds a "remarks" header in ``remarks_new_col``
    remarks_new_col: int | None = None
    # Channel-link columns (existing, or appended after the last used column), per platform.
    links: dict[str, int] | None = None
    # header cells the exporter must create: {"<column index>": "<header text>"}
    new_headers: dict[str, str] | None = None
    filled: list[str] | None = None  # id columns that contain at least one id (possible inputs)
    source: str | None = None  # chosen when the job starts
    targets: list[str] | None = None  # destination platforms (their id columns may be appended)

    def source_col(self, platform: str) -> int:
        return self.ids[platform]

    def dest_col(self, target: str) -> int:
        return self.ids[target]

    @property
    def remarks_col(self) -> int:
        col = self.remarks if self.remarks is not None else self.remarks_new_col
        assert col is not None
        return col

    def link_col(self, platform: str) -> int | None:
        return (self.links or {}).get(platform)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict[str, Any]) -> ColumnMap:
        return cls(**{k: d.get(k) for k in cls.__dataclass_fields__})  # type: ignore[arg-type]


@dataclass
class WorkbookAnalysis:
    sheet_name: str
    header_row: int
    headers: list[str]
    columns: ColumnMap
    detected_platform: str | None
    ambiguous: bool
    detection_reason: str
    row_count: int
    data_rows: int
    platforms: list[str] = field(default_factory=list)  # platforms with an id column, left to right
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        d = asdict(self)
        d["columns"] = self.columns.to_dict()
        return d


@dataclass
class ImportedRow:
    original_row: int
    source_value: str | None
    country: str | None
    existing_destinations: dict[str, str | None]  # destination platform -> current cell text
    existing_remarks: str | None


def cell_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value)
    return text if text.strip() else None


def _open(path: Path, **kwargs: Any) -> Any:
    if path.suffix.lower() == ".xls":
        raise WorkbookValidationError("Legacy .xls files are not supported. Please save the file as .xlsx.")
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise WorkbookValidationError(
            f"Unsupported file type '{path.suffix}'. Upload an .xlsx or .xlsm workbook."
        )
    if not zipfile.is_zipfile(path):
        raise WorkbookValidationError("The file is not a valid Excel workbook (corrupted or wrong format).")
    try:
        return load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm", **kwargs)
    except Exception as exc:
        raise WorkbookValidationError(
            f"The workbook could not be opened: {type(exc).__name__}: {exc}"
        ) from exc


def analyze_workbook(
    path: Path, source: str | None = None, targets: list[str] | None = None
) -> WorkbookAnalysis:
    """Locate the header row (>= 1 ``id_<platform>`` column) and detect the source platform.

    With ``source`` + ``targets`` (the user's choice when a job starts) the output layout is
    also planned: id columns for destinations the sheet lacks, a remarks column if missing,
    and one link column per platform are appended after the last used column.
    """
    wb = _open(path, data_only=True)
    if not wb.worksheets:
        raise WorkbookValidationError("The workbook contains no worksheets.")
    known = {"country", "remarks", *ID_HEADERS, *LINK_HEADERS.values()}
    matches: list[tuple[Any, int, dict[str, int], list[str]]] = []
    for ws in wb.worksheets:
        if getattr(ws, "sheet_state", "visible") != "visible" and len(wb.worksheets) > 1:
            continue
        for r in range(1, min(ws.max_row, HEADER_SCAN_ROWS) + 1):
            raw = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            cols: dict[str, int] = {}
            for idx, h in enumerate((normalize_header(v) for v in raw), start=1):
                if h in known and h not in cols:
                    cols[h] = idx
            if any(h in ID_HEADERS for h in cols):
                matches.append((ws, r, cols, [str(v) if v is not None else "" for v in raw]))
                break
    if not matches:
        raise WorkbookValidationError(
            "Required headers not found. Expected an ID column such as 'id_twitch' (supported: "
            f"{', '.join(ID_HEADERS)}), e.g. 'id_twitch | country | id_kick | id_youtube | remarks', "
            f"in the first {HEADER_SCAN_ROWS} rows."
        )
    warnings: list[str] = []
    if len(matches) > 1:
        warnings.append(
            f"Several sheets contain the expected headers; using '{matches[0][0].title}'. "
            "Other sheets are preserved unchanged."
        )
    ws, header_row, cols, headers = matches[0]
    ids = {ID_HEADERS[h]: c for h, c in sorted(cols.items(), key=lambda kv: kv[1]) if h in ID_HEADERS}
    platforms = list(ids)
    if source is not None:
        if source not in ids:
            raise WorkbookValidationError(
                f"The sheet has no id_{source} column; choose the source from: {', '.join(platforms)}."
            )
        bad = [t for t in targets or [] if t == source or t not in PLATFORMS]
        if not targets or bad:
            raise WorkbookValidationError(
                "Choose at least one destination platform, different from the source "
                f"({', '.join(p for p in PLATFORMS if p != source)})."
            )
    if "country" not in cols:
        warnings.append("No 'country' column found; country evidence will not be used.")

    last = _last_data_row(ws, header_row, list(cols.values()))
    data_rows = max(0, last - header_row)
    if data_rows > MAX_DATA_ROWS:
        raise WorkbookValidationError(f"Too many rows ({data_rows}); the limit is {MAX_DATA_ROWS}.")
    fill = {
        p: sum(1 for r in range(header_row + 1, last + 1) if not is_blank(ws.cell(row=r, column=c).value))
        for p, c in ids.items()
    }
    if data_rows == 0 or not any(fill.values()):
        raise WorkbookValidationError("The workbook has the expected headers but no data rows.")

    detected: str | None
    filled = [p for p in platforms if fill[p]]  # only these can be the input
    if source is not None and source not in filled:
        raise WorkbookValidationError(
            f"The {header_of(source)} column is empty; the source must be a column that has IDs "
            f"({', '.join(header_of(p) for p in filled)})."
        )
    rates = ", ".join(f"{header_of(p)} {fill[p] / data_rows:.0%}" for p in platforms)
    if source is not None:
        detected, ambiguous = source, False
        reason = f"Source {label(source)}; searching {', '.join(label(t) for t in targets or [])}."
    elif len(filled) == 1:
        detected, ambiguous = filled[0], False
        reason = f"'{header_of(filled[0])}' is the only ID column with IDs in it (fill: {rates})."
    else:
        first = filled[0]
        fuller = [p for p in filled[1:] if fill[p] > fill[first] + 0.2 * data_rows]
        if not fuller:
            detected, ambiguous = first, False
            reason = f"'{header_of(first)}' is the first ID column with IDs (fill: {rates})."
        else:
            detected, ambiguous = None, True
            reason = (
                f"Column order suggests a {label(first).upper()} source, but '{header_of(fuller[0])}' is "
                f"fuller (fill: {rates}). Please choose the source platform."
            )

    remarks_new_col = None
    links: dict[str, int] = {}
    new_headers: dict[str, str] = {}
    if source is not None:
        # plan the appended columns: missing destination ids, remarks, then link columns
        next_col = 1 + max([c for c in range(1, ws.max_column + 1) if _column_has_data(ws, c)] + [0])
        for t in targets or []:
            if t not in ids:
                ids[t] = next_col
                new_headers[str(next_col)] = header_of(t)
                warnings.append(
                    f"No '{header_of(t)}' column; it will be added in column {get_column_letter(next_col)}."
                )
                next_col += 1
        if "remarks" not in cols:
            remarks_new_col = next_col
            warnings.append(
                f"No 'remarks' column; one will be added in column {get_column_letter(next_col)}."
            )
            next_col += 1
        for p in sorted([source, *(targets or [])], key=lambda p: ids[p]):
            header = LINK_HEADERS[p]
            if header in cols:
                links[p] = cols[header]
            else:
                links[p] = next_col
                new_headers[str(next_col)] = header
                next_col += 1
    colmap = ColumnMap(
        sheet_name=ws.title,
        header_row=header_row,
        ids=ids,
        country=cols.get("country"),
        remarks=cols.get("remarks"),
        remarks_new_col=remarks_new_col,
        links=links or None,
        new_headers=new_headers or None,
        filled=filled,
        source=source,
        targets=list(targets) if targets else None,
    )

    src_col = ids[detected or platforms[0]]
    for r in range(header_row + 1, last + 1):
        src = ws.cell(row=r, column=src_col).value
        if isinstance(src, str) and (" " in src.strip() or len(src.strip()) > 64):
            warnings.append(f"Row {r}: source value {src!r} looks malformed (spaces/too long).")
            if len(warnings) > 25:
                warnings.append("…more warnings suppressed.")
                break
    return WorkbookAnalysis(
        sheet_name=ws.title,
        header_row=header_row,
        headers=headers,
        columns=colmap,
        detected_platform=detected,
        ambiguous=ambiguous,
        detection_reason=reason,
        row_count=data_rows,
        data_rows=data_rows,
        platforms=platforms,
        warnings=warnings,
    )


def header_of(platform: str) -> str:
    return f"id_{platform}"


def _last_data_row(ws: Any, header_row: int, cols: list[int]) -> int:
    last = header_row
    for r in range(ws.max_row, header_row, -1):
        if any(not is_blank(ws.cell(row=r, column=c).value) for c in cols):
            last = r
            break
    return last


def read_rows(path: Path, columns: ColumnMap) -> list[ImportedRow]:
    assert columns.source and columns.targets, "plan the columns with analyze_workbook(source, targets)"
    wb = _open(path, data_only=True)
    ws = wb[columns.sheet_name]
    src_col = columns.source_col(columns.source)
    dest = {t: columns.dest_col(t) for t in columns.targets}
    cols = [src_col, *dest.values()] + [c for c in (columns.country, columns.remarks) if c]
    last = _last_data_row(ws, columns.header_row, cols)
    rows: list[ImportedRow] = []
    for r in range(columns.header_row + 1, last + 1):
        rows.append(
            ImportedRow(
                original_row=r,
                source_value=cell_text(ws.cell(row=r, column=src_col).value),
                country=cell_text(ws.cell(row=r, column=columns.country).value) if columns.country else None,
                existing_destinations={t: cell_text(ws.cell(row=r, column=c).value) for t, c in dest.items()},
                existing_remarks=cell_text(ws.cell(row=r, column=columns.remarks).value)
                if columns.remarks
                else None,
            )
        )
    return rows


def _column_has_data(ws: Any, col: int) -> bool:
    return any(ws.cell(row=r, column=col).value not in (None, "") for r in range(1, ws.max_row + 1))
